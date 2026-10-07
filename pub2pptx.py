#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.12"
# dependencies = [
#   "python-pptx>=1.0",
#   "pillow>=10.1",
# ]
# ///
"""
pub2pptx - convert Microsoft Publisher (.pub) files to PowerPoint (.pptx).

The Publisher file is read with `pub2raw` from libmspub and every page is rebuilt as
a native slide: pictures, shapes, text boxes (with fonts, sizes, colours, alignment)
and tables become real, editable PowerPoint objects. Each file also gets a first-page
preview image, painted with Pillow, that file managers show as its thumbnail.

`--mode snapshot` instead paints every page with Pillow and places the picture on a
slide: not editable, but a single object per page.

External tool:
  pub2raw     Arch: `pacman -S libmspub`   Debian/Ubuntu: `apt install libmspub-tools`
              macOS: `brew install libmspub`
              Windows: MSYS2 `pacman -S mingw-w64-ucrt-x86_64-libmspub` (found automatically
              under C:\\msys64, or point the PUB2RAW environment variable at pub2raw.exe)

Examples:
  uv run pub2pptx.py brochure.pub                  # -> brochure.pptx
  uv run pub2pptx.py brochure.pub -o out/deck.pptx
  uv run pub2pptx.py *.pub -o converted/           # many files into a folder
  uv run pub2pptx.py menu.pub --font-map "Raleway=Calibri,Poppins=Segoe UI"
  uv run pub2pptx.py menu.pub --mode snapshot --dpi 200
"""
import argparse
import base64
import hashlib
import io
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import traceback
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Pt

VERBOSE = False
EMU_PER_INCH = 914400


def log(msg: str) -> None:
    print(msg, file=sys.stderr)


def vlog(msg: str) -> None:
    if VERBOSE:
        print(msg, file=sys.stderr)


# --------------------------------------------------------------------------- #
# Parsing the pub2raw callback stream
# --------------------------------------------------------------------------- #

_LINE_RE = re.compile(r"^\s*([A-Za-z]\w*)\s*(?:\((.*)\))?\s*$")
# Structural tokens inside a property list: parens, and the ", " that separates
# two key/value pairs (identified by a following `prefix:name: ` key).
_TOK_RE = re.compile(r"\(|\)|, (?=[A-Za-z]+:[A-Za-z0-9-]+: )")


@dataclass(slots=True)
class Event:
    name: str
    props: dict[str, Any]
    text: str = ""  # only for insertText


def parse_props(s: str) -> dict[str, Any]:
    """Parse `k: v, k: v, k: ((k: v), (k: v))` into a dict (lists -> list[dict])."""
    s = s.strip()
    if not s:
        return {}
    parts: list[str] = []
    depth = 0
    start = 0
    for m in _TOK_RE.finditer(s):
        tok = m.group()
        if tok == "(":
            depth += 1
        elif tok == ")":
            depth -= 1
        elif depth == 0:
            parts.append(s[start : m.start()])
            start = m.end()
    parts.append(s[start:])
    out: dict[str, Any] = {}
    for part in parts:
        key, sep, val = part.partition(": ")
        if not sep:
            continue
        val = val.strip()
        out[key.strip()] = parse_list(val) if val.startswith("(") else val
    return out


def parse_list(v: str) -> list[dict[str, Any]]:
    """Parse `((k: v, ...), (k: v, ...))` into a list of dicts."""
    inner = v.strip()
    if inner.startswith("(") and inner.endswith(")"):
        inner = inner[1:-1]
    items: list[dict[str, Any]] = []
    depth = 0
    start = -1
    for i, ch in enumerate(inner):
        if ch == "(":
            if depth == 0:
                start = i + 1
            depth += 1
        elif ch == ")":
            depth -= 1
            if depth == 0 and start >= 0:
                items.append(parse_props(inner[start:i]))
                start = -1
    return items


_TEXT_OPEN_RE = re.compile(r"^\s*insertText \((.*)$", re.S)


def parse_raw(text: str) -> list[Event]:
    events: list[Event] = []
    pending: list[str] | None = None  # lines of an insertText whose text spans several lines
    # Split on "\n" only: Publisher text can contain "\r" (a line break) inside a run.
    for line in text.split("\n"):
        if pending is not None:
            pending.append(line)
            if line.rstrip().endswith(")"):
                joined = "\n".join(pending).rstrip()[:-1]
                events.append(Event("insertText", {}, joined))
                pending = None
            continue
        if not line.strip():
            continue
        m = _LINE_RE.match(line)
        if not m:
            mo = _TEXT_OPEN_RE.match(line)
            if mo:
                pending = [mo.group(1)]
            else:
                vlog(f"skipping unparsable line: {line[:80]!r}")
            continue
        name, args = m.group(1), m.group(2)
        if name == "insertText":
            events.append(Event(name, {}, args or ""))
        else:
            events.append(Event(name, parse_props(args or "")))
    if pending is not None:
        events.append(Event("insertText", {}, "\n".join(pending)))
    return events


# --------------------------------------------------------------------------- #
# Units and small helpers
# --------------------------------------------------------------------------- #

_LEN_RE = re.compile(r"^\s*(-?\d*\.?\d+(?:[eE][-+]?\d+)?)\s*(in|pt|cm|mm|px|pc|%)?\s*$")
_TO_INCH = {"in": 1.0, "pt": 1 / 72, "cm": 1 / 2.54, "mm": 1 / 25.4, "px": 1 / 96, "pc": 1 / 6, None: 1.0}


def length(v: Any, default: float | None = None) -> float | None:
    """Convert a librevenge length ('1.5in', '12pt') to inches."""
    if v is None:
        return default
    m = _LEN_RE.match(str(v))
    if not m or m.group(2) == "%":
        return default
    return float(m.group(1)) * _TO_INCH[m.group(2)]


def percent(v: Any, default: float | None = None) -> float | None:
    """'50%' -> 0.5, '0.5' -> 0.5"""
    if v is None:
        return default
    s = str(v).strip()
    try:
        return float(s[:-1]) / 100 if s.endswith("%") else float(s)
    except ValueError:
        return default


def emu(inches: float) -> int:
    return int(round(inches * EMU_PER_INCH))


def rgb(v: Any) -> RGBColor | None:
    if not v:
        return None
    s = str(v).strip()
    if s.startswith("#") and len(s) == 7:
        try:
            return RGBColor.from_string(s[1:])
        except ValueError:
            return None
    return None


def as_float(v: Any, default: float = 0.0) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


# --------------------------------------------------------------------------- #
# Document model
# --------------------------------------------------------------------------- #

@dataclass(slots=True)
class Run:
    props: dict[str, Any]
    text: str = ""  # may contain "\t" and "\n" (line break)


@dataclass(slots=True)
class Paragraph:
    props: dict[str, Any]
    runs: list[Run] = field(default_factory=list)
    bullet: str | None = None


@dataclass(slots=True)
class Cell:
    row: int
    col: int
    props: dict[str, Any]
    paragraphs: list[Paragraph] = field(default_factory=list)


@dataclass(slots=True)
class Row:
    props: dict[str, Any]
    cells: list[Cell] = field(default_factory=list)


@dataclass(slots=True)
class TextBox:
    props: dict[str, Any]
    style: dict[str, Any]
    paragraphs: list[Paragraph] = field(default_factory=list)


@dataclass(slots=True)
class Table:
    props: dict[str, Any]
    style: dict[str, Any]
    rows: list[Row] = field(default_factory=list)


@dataclass(slots=True)
class Shape:
    kind: str  # polygon, polyline, path, rectangle, ellipse, graphic, connector
    props: dict[str, Any]
    style: dict[str, Any]


type Item = Shape | TextBox | Table


@dataclass(slots=True)
class Page:
    props: dict[str, Any]
    items: list[Item] = field(default_factory=list)


@dataclass(slots=True)
class Document:
    meta: dict[str, Any] = field(default_factory=dict)
    pages: list[Page] = field(default_factory=list)


def build_document(events: list[Event]) -> Document:
    doc = Document()
    page: Page | None = None
    style: dict[str, Any] = {}
    # Text container state
    paragraphs: list[Paragraph] | None = None  # current paragraph list target
    para: Paragraph | None = None
    span_props: dict[str, Any] = {}
    table: Table | None = None
    row: Row | None = None
    list_levels: list[dict[str, Any]] = []

    def add_run(text: str) -> None:
        nonlocal para
        if para is None:
            # Text outside a paragraph: open an implicit one.
            para = Paragraph({})
            if paragraphs is not None:
                paragraphs.append(para)
        if para.runs and para.runs[-1].props is span_props:
            para.runs[-1].text += text
        else:
            para.runs.append(Run(span_props, text))

    for ev in events:
        p = ev.props
        match ev.name:
            case "setDocumentMetaData":
                doc.meta = p
            case "startPage":
                page = Page(p)
                doc.pages.append(page)
            case "endPage":
                page = None
            case "setStyle":
                style = p
            case ("drawPolygon" | "drawPolyline" | "drawPath" | "drawRectangle" | "drawEllipse"
                  | "drawGraphicObject" | "drawConnector") as name:
                if page is not None:
                    page.items.append(Shape(name.removeprefix("draw").lower(), p, style))
            case "startTextObject":
                tb = TextBox(p, style)
                if page is not None:
                    page.items.append(tb)
                paragraphs = tb.paragraphs
            case "endTextObject":
                paragraphs = None
                para = None
            case "startTableObject":
                table = Table(p, style)
                if page is not None:
                    page.items.append(table)
            case "endTableObject":
                table = None
                paragraphs = None
            case "openTableRow":
                if table is not None:
                    row = Row(p)
                    table.rows.append(row)
            case "closeTableRow":
                row = None
            case "openTableCell":
                if row is not None and table is not None:
                    cell = Cell(int(as_float(p.get("librevenge:row"), len(table.rows) - 1)),
                                int(as_float(p.get("librevenge:column"), len(row.cells))), p)
                    row.cells.append(cell)
                    paragraphs = cell.paragraphs
            case "closeTableCell":
                paragraphs = None
                para = None
            case "openOrderedListLevel" | "openUnorderedListLevel":
                list_levels.append(p)
            case "closeOrderedListLevel" | "closeUnorderedListLevel":
                if list_levels:
                    list_levels.pop()
            case "openParagraph" | "openListElement" as name:
                para = Paragraph(p)
                if name == "openListElement":
                    lvl = list_levels[-1] if list_levels else {}
                    para.bullet = lvl.get("text:bullet-char") or ("1." if "style:num-format" in lvl else "\u2022")
                    para.props = {**lvl, **p}
                if paragraphs is not None:
                    paragraphs.append(para)
            case "closeParagraph" | "closeListElement":
                # A trailing "\r" is Publisher's paragraph mark, not an extra line break.
                if para is not None and para.runs and para.runs[-1].text.endswith("\n"):
                    para.runs[-1].text = para.runs[-1].text[:-1]
                para = None
            case "openSpan":
                span_props = p
            case "closeSpan":
                span_props = {}
            case "insertText":
                add_run(ev.text.replace("\r\n", "\n").replace("\r", "\n"))
            case "insertTab":
                add_run("\t")
            case "insertSpace":
                add_run(" ")
            case "insertLineBreak":
                add_run("\n")
            case "insertField":
                add_run(p.get("librevenge:field-type", ""))
            case _:
                pass  # layers, groups, covered cells, ... carry nothing we need
    return doc


# --------------------------------------------------------------------------- #
# Geometry helpers
# --------------------------------------------------------------------------- #

def points_of(plist: Any) -> list[tuple[float, float]]:
    pts = []
    for d in plist or []:
        x, y = length(d.get("svg:x")), length(d.get("svg:y"))
        if x is not None and y is not None:
            pts.append((x, y))
    return pts


def as_rectangle(pts: list[tuple[float, float]]):
    """If pts describe a (possibly rotated) rectangle, return (cx, cy, w, h, angle_deg)."""
    if len(pts) == 5 and abs(pts[0][0] - pts[4][0]) < 1e-6 and abs(pts[0][1] - pts[4][1]) < 1e-6:
        pts = pts[:4]
    if len(pts) != 4:
        return None
    (x0, y0), (x1, y1), (x2, y2), (x3, y3) = pts
    e0 = (x1 - x0, y1 - y0)
    e1 = (x2 - x1, y2 - y1)
    e2 = (x3 - x2, y3 - y2)
    e3 = (x0 - x3, y0 - y3)
    w = math.hypot(*e0)
    h = math.hypot(*e1)
    if w < 1e-6 or h < 1e-6:
        return None
    tol = 1e-3 * max(w, h)
    # opposite sides equal and parallel, adjacent sides perpendicular
    if abs(e0[0] + e2[0]) > tol or abs(e0[1] + e2[1]) > tol:
        return None
    if abs(e1[0] + e3[0]) > tol or abs(e1[1] + e3[1]) > tol:
        return None
    if abs(e0[0] * e1[0] + e0[1] * e1[1]) > tol * max(w, h):
        return None
    angle = math.degrees(math.atan2(e0[1], e0[0]))
    cx = (x0 + x1 + x2 + x3) / 4
    cy = (y0 + y1 + y2 + y3) / 4
    return cx, cy, w, h, angle


def arc_points(x1, y1, rx, ry, phi_deg, large_arc, sweep, x2, y2, n=24):
    """Flatten an SVG elliptical arc into points (excluding the start point)."""
    if rx == 0 or ry == 0 or (x1 == x2 and y1 == y2):
        return [(x2, y2)]
    phi = math.radians(phi_deg)
    cp, sp = math.cos(phi), math.sin(phi)
    dx, dy = (x1 - x2) / 2, (y1 - y2) / 2
    x1p = cp * dx + sp * dy
    y1p = -sp * dx + cp * dy
    rx, ry = abs(rx), abs(ry)
    lam = (x1p ** 2) / (rx ** 2) + (y1p ** 2) / (ry ** 2)
    if lam > 1:
        rx *= math.sqrt(lam)
        ry *= math.sqrt(lam)
    num = rx ** 2 * ry ** 2 - rx ** 2 * y1p ** 2 - ry ** 2 * x1p ** 2
    den = rx ** 2 * y1p ** 2 + ry ** 2 * x1p ** 2
    coef = math.sqrt(max(0.0, num / den)) if den else 0.0
    if large_arc == sweep:
        coef = -coef
    cxp = coef * rx * y1p / ry
    cyp = -coef * ry * x1p / rx
    cx = cp * cxp - sp * cyp + (x1 + x2) / 2
    cy = sp * cxp + cp * cyp + (y1 + y2) / 2

    def ang(ux, uy, vx, vy):
        a = math.atan2(ux * vy - uy * vx, ux * vx + uy * vy)
        return a

    t1 = ang(1, 0, (x1p - cxp) / rx, (y1p - cyp) / ry)
    dt = ang((x1p - cxp) / rx, (y1p - cyp) / ry, (-x1p - cxp) / rx, (-y1p - cyp) / ry)
    if not sweep and dt > 0:
        dt -= 2 * math.pi
    elif sweep and dt < 0:
        dt += 2 * math.pi
    pts = []
    for i in range(1, n + 1):
        t = t1 + dt * i / n
        ex, ey = rx * math.cos(t), ry * math.sin(t)
        pts.append((cp * ex - sp * ey + cx, sp * ex + cp * ey + cy))
    pts[-1] = (x2, y2)
    return pts


def path_commands(d: list[dict[str, Any]]):
    """
    Turn a librevenge svg:d list into a list of commands:
      ("M", x, y) ("L", x, y) ("C", x1, y1, x2, y2, x, y) ("Q", x1, y1, x, y) ("Z",)
    Arcs are flattened into line segments.
    """
    cmds: list[tuple] = []
    cur = (0.0, 0.0)
    start = (0.0, 0.0)
    last_ctrl: tuple[float, float] | None = None
    last_cmd = ""
    for seg in d or []:
        a = str(seg.get("librevenge:path-action", "")).upper()
        gx = length(seg.get("svg:x"), cur[0])
        gy = length(seg.get("svg:y"), cur[1])
        if a == "M":
            cur = start = (gx, gy)
            cmds.append(("M", gx, gy))
            last_ctrl = None
        elif a == "L":
            cur = (gx, gy)
            cmds.append(("L", gx, gy))
            last_ctrl = None
        elif a == "H":
            cur = (gx, cur[1])
            cmds.append(("L", *cur))
            last_ctrl = None
        elif a == "V":
            cur = (cur[0], gy)
            cmds.append(("L", *cur))
            last_ctrl = None
        elif a in ("C", "S"):
            if a == "C":
                x1, y1 = length(seg.get("svg:x1"), cur[0]), length(seg.get("svg:y1"), cur[1])
            else:
                x1, y1 = (2 * cur[0] - last_ctrl[0], 2 * cur[1] - last_ctrl[1]) if (last_ctrl and last_cmd in "CS") else cur
            x2, y2 = length(seg.get("svg:x2"), cur[0]), length(seg.get("svg:y2"), cur[1])
            cmds.append(("C", x1, y1, x2, y2, gx, gy))
            last_ctrl = (x2, y2)
            cur = (gx, gy)
        elif a in ("Q", "T"):
            if a == "Q":
                x1, y1 = length(seg.get("svg:x1"), cur[0]), length(seg.get("svg:y1"), cur[1])
            else:
                x1, y1 = (2 * cur[0] - last_ctrl[0], 2 * cur[1] - last_ctrl[1]) if (last_ctrl and last_cmd in "QT") else cur
            cmds.append(("Q", x1, y1, gx, gy))
            last_ctrl = (x1, y1)
            cur = (gx, gy)
        elif a == "A":
            rx, ry = length(seg.get("svg:rx"), 0.0), length(seg.get("svg:ry"), 0.0)
            rot = as_float(seg.get("librevenge:rotate"), 0.0)
            large = str(seg.get("librevenge:large-arc", "false")).lower() in ("true", "1")
            sweep = str(seg.get("librevenge:sweep", "false")).lower() in ("true", "1")
            for px, py in arc_points(cur[0], cur[1], rx, ry, rot, large, sweep, gx, gy):
                cmds.append(("L", px, py))
            cur = (gx, gy)
            last_ctrl = None
        elif a == "Z":
            cmds.append(("Z",))
            cur = start
            last_ctrl = None
        last_cmd = a
    return cmds


def commands_bbox(cmds) -> tuple[float, float, float, float] | None:
    xs, ys = [], []
    for c in cmds:
        for i in range(1, len(c), 2):
            xs.append(c[i])
            ys.append(c[i + 1])
    if not xs:
        return None
    return min(xs), min(ys), max(xs), max(ys)


# --------------------------------------------------------------------------- #
# PPTX rendering
# --------------------------------------------------------------------------- #

NO_STYLE_NO_GRID = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
_ALIGN = {
    "left": PP_ALIGN.LEFT, "start": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY,
}
_VALIGN = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


class Renderer:
    def __init__(self, doc: Document, font_map: dict[str, str] | None = None, blank_size: float | None = None):
        self.doc = doc
        self.font_map = font_map or {}
        self.blank_size = blank_size  # forced size (pt) for blank table-cell lines, or None to estimate
        self.prs = Presentation()
        self._image_cache: dict[str, bytes | None] = {}
        self._native_parts: dict[str, Any] = {}
        self.warnings: list[str] = []
        first = doc.pages[0].props if doc.pages else {}
        self.page_w = length(first.get("svg:width"), 11.6929) or 11.6929
        self.page_h = length(first.get("svg:height"), 8.2677) or 8.2677
      
        # PowerPoint slide dimension limits (in inches)
        MIN_INCHES = 1.0
        MAX_INCHES = 56.0

        # Scale down if too large
        if self.page_w > MAX_INCHES or self.page_h > MAX_INCHES:
            scale = min(MAX_INCHES / self.page_w, MAX_INCHES / self.page_h)
            self.page_w *= scale
            self.page_h *= scale

        # Scale up if too small
        if self.page_w < MIN_INCHES or self.page_h < MIN_INCHES:
            scale = max(MIN_INCHES / self.page_w, MIN_INCHES / self.page_h)
            self.page_w *= scale
            self.page_h *= scale
      
        self.prs.slide_width = emu(self.page_w)
        self.prs.slide_height = emu(self.page_h)
        self._blank = self.prs.slide_layouts[6]

    # ---- top level -------------------------------------------------------- #
    def render(self) -> Presentation:
        for i, page in enumerate(self.doc.pages, 1):
            slide = self.prs.slides.add_slide(self._blank)
            vlog(f"page {i}: {len(page.items)} objects")
            for item in page.items:
                try:
                    if isinstance(item, Shape):
                        self.render_shape(slide, item)
                    elif isinstance(item, TextBox):
                        self.render_textbox(slide, item)
                    elif isinstance(item, Table):
                        self.render_table(slide, item)
                except Exception as e:  # keep going, one bad object shouldn't sink the page
                    self.warn(f"page {i}: could not convert {type(item).__name__}: {e}")
                    if VERBOSE:
                        traceback.print_exc()
        return self.prs

    def warn(self, msg: str) -> None:
        self.warnings.append(msg)
        log(f"warning: {msg}")

    # ---- images ----------------------------------------------------------- #
    NATIVE_ONLY = {  # formats Pillow can't decode but PowerPoint renders natively
        "image/x-wmf": ("wmf", CT.X_WMF), "image/wmf": ("wmf", CT.X_WMF),
        "image/x-emf": ("emf", CT.X_EMF), "image/emf": ("emf", CT.X_EMF),
    }

    def image_bytes(self, b64: str, mime: str) -> bytes | None:
        """Decode an embedded image (base64) to bytes."""
        key = b64[:64] + str(len(b64))
        if key in self._image_cache:
            return self._image_cache[key]
        try:
            data = base64.b64decode(b64)
        except Exception:
            data = None
        self._image_cache[key] = data
        return data

    def image_rId(self, slide, data: bytes, mime: str) -> str | None:
        """
        Add an image part to the slide and return its relationship id.

        python-pptx sniffs images with Pillow, which cannot read WMF/EMF. Those are
        embedded as raw parts with the right content type instead: PowerPoint draws
        them itself. Anything else unreadable is skipped with a warning.
        """
        try:
            _part, rId = slide.part.get_or_add_image_part(io.BytesIO(data))
            return rId
        except Exception:
            pass
        native = self.NATIVE_ONLY.get(mime.lower())
        if native is None:
            self.warn(f"skipping image of unsupported type {mime or 'unknown'}")
            return None
        ext, content_type = native
        from pptx.parts.image import ImagePart
        digest = ext + hashlib.sha1(data).hexdigest()
        part = self._native_parts.get(digest)
        if part is None:
            package = slide.part.package
            partname = package.next_partname(f"/ppt/media/image%d.{ext}")
            part = ImagePart(partname, content_type, package, data)
            self._native_parts[digest] = part
        return slide.part.relate_to(part, RT.IMAGE)

    def add_picture(self, slide, data: bytes, mime: str, x: float, y: float, w: float, h: float):
        """A Picture shape from image bytes at (x, y, w, h) inches; None if the image is unusable."""
        try:
            return slide.shapes.add_picture(io.BytesIO(data), emu(x), emu(y), emu(w), emu(h))
        except Exception:
            pass
        rId = self.image_rId(slide, data, mime)
        if rId is None:
            return None
        from pptx.oxml.shapes.picture import CT_Picture
        shapes = slide.shapes
        shape_id = shapes._next_shape_id
        pic = CT_Picture.new_pic(shape_id, f"Picture {shape_id - 1}", "image", rId, emu(x), emu(y), emu(w), emu(h))
        shapes._spTree.append(pic)
        return shapes[-1]

    def add_blip_fill(self, slide, shape, data: bytes, mime: str, tile: bool) -> None:
        """Give an autoshape a picture fill (tiled or stretched)."""
        rId = self.image_rId(slide, data, mime)
        if rId is None:
            shape.fill.background()
            return
        spPr = shape._element.spPr
        for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill"):
            for el in spPr.findall(qn(tag)):
                spPr.remove(el)
        blipFill = spPr.makeelement(qn("a:blipFill"), {})
        blip = blipFill.makeelement(qn("a:blip"), {qn("r:embed"): rId})
        blipFill.append(blip)
        if tile:
            blipFill.append(blipFill.makeelement(
                qn("a:tile"), {"tx": "0", "ty": "0", "sx": "100000", "sy": "100000", "flip": "none", "algn": "tl"}))
        else:
            stretch = blipFill.makeelement(qn("a:stretch"), {})
            stretch.append(stretch.makeelement(qn("a:fillRect"), {}))
            blipFill.append(stretch)
        geom = spPr.find(qn("a:prstGeom"))
        if geom is None:
            geom = spPr.find(qn("a:custGeom"))
        if geom is not None:
            geom.addnext(blipFill)
        else:
            spPr.append(blipFill)

    # ---- fill / stroke ---------------------------------------------------- #
    def apply_style(self, slide, shape, style: dict[str, Any], *, allow_fill=True) -> None:
        fill = style.get("draw:fill", "none")
        if allow_fill:
            if fill == "solid" and rgb(style.get("draw:fill-color")):
                shape.fill.solid()
                shape.fill.fore_color.rgb = rgb(style.get("draw:fill-color"))
                op = percent(style.get("draw:opacity"))
                if op is not None and op < 1:
                    clr = shape.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
                    clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(op * 100000))}))
            elif fill == "gradient":
                stops = style.get("svg:linearGradient") or style.get("svg:radialGradient") or []
                colors = [rgb(s.get("svg:stop-color")) for s in stops]
                colors = [c for c in colors if c] or [rgb(style.get("draw:start-color")), rgb(style.get("draw:end-color"))]
                colors = [c for c in colors if c]
                if colors:
                    shape.fill.gradient()
                    gs = shape.fill.gradient_stops
                    gs[0].color.rgb = colors[0]
                    gs[1].color.rgb = colors[-1]
                    angle = as_float(style.get("draw:angle"), 0.0)
                    try:
                        shape.fill.gradient_angle = (90 - angle) % 360
                    except Exception:
                        pass
            elif fill == "bitmap" and style.get("draw:fill-image"):
                data = self.image_bytes(style["draw:fill-image"], style.get("librevenge:mime-type", ""))
                if data:
                    tile = style.get("style:repeat", "repeat") not in ("stretch", "no-repeat")
                    self.add_blip_fill(slide, shape, data, style.get("librevenge:mime-type", ""), tile)
                else:
                    shape.fill.background()
            else:
                shape.fill.background()

        stroke = style.get("draw:stroke", "none")
        if stroke == "none":
            shape.line.fill.background()
        else:
            color = rgb(style.get("svg:stroke-color")) or RGBColor(0, 0, 0)
            shape.line.color.rgb = color
            w = length(style.get("svg:stroke-width"))
            if w is not None:
                shape.line.width = emu(max(w, 0.001))
            if stroke == "dash":
                shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH

    # ---- shapes ----------------------------------------------------------- #
    def render_shape(self, slide, item: Shape) -> None:
        st, p = item.style, item.props
        rot = as_float(p.get("librevenge:rotate", st.get("librevenge:rotate")), 0.0)
        match item.kind:
            case "graphicobject":
                data = self.image_bytes(p.get("office:binary-data", ""), p.get("librevenge:mime-type", ""))
                if not data:
                    return
                x, y = length(p.get("svg:x"), 0.0), length(p.get("svg:y"), 0.0)
                w, h = length(p.get("svg:width"), 1.0), length(p.get("svg:height"), 1.0)
                pic = self.add_picture(slide, data, p.get("librevenge:mime-type", ""), x, y, w, h)
                if pic is not None and rot:
                    pic.rotation = rot
            case "polygon":
                pts = points_of(p.get("svg:points"))
                if rect := as_rectangle(pts):
                    cx, cy, w, h, angle = rect
                    self._rect_like(slide, st, cx - w / 2, cy - h / 2, w, h, angle)
                elif pts:
                    self._freeform(slide, st, [("M", *pts[0]), *(("L", *q) for q in pts[1:]), ("Z",)], rot)
            case "polyline" | "connector":
                pts = points_of(p.get("svg:points"))
                if len(pts) >= 2:
                    self._freeform(slide, st, [("M", *pts[0]), *(("L", *q) for q in pts[1:])], rot, closed=False)
            case "rectangle":
                x, y = length(p.get("svg:x"), 0.0), length(p.get("svg:y"), 0.0)
                w, h = length(p.get("svg:width"), 0.0), length(p.get("svg:height"), 0.0)
                self._rect_like(slide, st, x, y, w, h, rot, rounded=length(p.get("svg:rx")))
            case "ellipse":
                cx, cy = length(p.get("svg:cx"), 0.0), length(p.get("svg:cy"), 0.0)
                rx, ry = length(p.get("svg:rx"), 0.0), length(p.get("svg:ry"), 0.0)
                shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, emu(cx - rx), emu(cy - ry), emu(2 * rx), emu(2 * ry))
                self.apply_style(slide, shp, st)
                if rot:
                    shp.rotation = rot
            case "path":
                cmds = path_commands(p.get("svg:d") or [])
                self._freeform(slide, st, cmds, rot, closed=any(c[0] == "Z" for c in cmds))

    def _rect_like(self, slide, st, x, y, w, h, angle, rounded=None) -> None:
        """A rectangle: becomes a Picture when it is a stretched bitmap, else an autoshape."""
        if w <= 0 or h <= 0:
            return
        is_stretch_bitmap = (st.get("draw:fill") == "bitmap" and st.get("style:repeat") == "stretch"
                             and st.get("draw:stroke", "none") == "none" and not rounded)
        if is_stretch_bitmap:
            data = self.image_bytes(st.get("draw:fill-image", ""), st.get("librevenge:mime-type", ""))
            if data:
                pic = self.add_picture(slide, data, st.get("librevenge:mime-type", ""), x, y, w, h)
                if pic is not None and abs(angle) > 1e-3:
                    pic.rotation = angle
            return
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
        shp = slide.shapes.add_shape(kind, emu(x), emu(y), emu(w), emu(h))
        if rounded and min(w, h) > 0:
            try:
                shp.adjustments[0] = min(0.5, rounded / min(w, h))
            except Exception:
                pass
        self.apply_style(slide, shp, st)
        if abs(angle) > 1e-3:
            shp.rotation = angle

    def _freeform(self, slide, st, cmds, rot=0.0, closed=True) -> None:
        bbox = commands_bbox(cmds)
        if not bbox:
            return
        x0, y0, x1, y1 = bbox
        w, h = max(x1 - x0, 1e-4), max(y1 - y0, 1e-4)
        W, H = emu(w), emu(h)
        builder = slide.shapes.build_freeform(0, 0, scale=1.0)
        builder.add_line_segments([(W, 0), (W, H), (0, H)], close=True)
        shp = builder.convert_to_shape(emu(x0), emu(y0))
        # Rewrite the path with the real geometry.
        pathLst = shp._element.spPr.find(qn("a:custGeom")).find(qn("a:pathLst"))
        for el in list(pathLst):
            pathLst.remove(el)
        path = pathLst.makeelement(qn("a:path"), {"w": str(W), "h": str(H)})
        if not closed and st.get("draw:fill", "none") == "none":
            path.set("fill", "none")
        pathLst.append(path)

        def pt(x, y):
            return {"x": str(int(round((x - x0) * EMU_PER_INCH))), "y": str(int(round((y - y0) * EMU_PER_INCH)))}

        def add(tag, *coords):
            el = path.makeelement(qn(tag), {})
            for i in range(0, len(coords), 2):
                el.append(el.makeelement(qn("a:pt"), pt(coords[i], coords[i + 1])))
            path.append(el)

        open_sub = False
        for c in cmds:
            if c[0] == "M":
                if open_sub and closed:
                    path.append(path.makeelement(qn("a:close"), {}))
                add("a:moveTo", c[1], c[2])
                open_sub = True
            elif c[0] == "L":
                add("a:lnTo", c[1], c[2])
            elif c[0] == "C":
                add("a:cubicBezTo", *c[1:])
            elif c[0] == "Q":
                add("a:quadBezTo", *c[1:])
            elif c[0] == "Z":
                path.append(path.makeelement(qn("a:close"), {}))
                open_sub = False
        if open_sub and closed:
            path.append(path.makeelement(qn("a:close"), {}))
        self.apply_style(slide, shp, st)
        if rot:
            shp.rotation = rot

    # ---- text ------------------------------------------------------------- #
    def render_textbox(self, slide, item: TextBox) -> None:
        p = item.props
        x, y = length(p.get("svg:x"), 0.0), length(p.get("svg:y"), 0.0)
        w, h = max(length(p.get("svg:width"), 1.0), 0.05), max(length(p.get("svg:height"), 0.5), 0.05)
        shp = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
        self.apply_style(slide, shp, item.style)
        tf = shp.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = emu(length(p.get("fo:padding-left"), 0.04))
        tf.margin_right = emu(length(p.get("fo:padding-right"), 0.04))
        tf.margin_top = emu(length(p.get("fo:padding-top"), 0.04))
        tf.margin_bottom = emu(length(p.get("fo:padding-bottom"), 0.04))
        tf.vertical_anchor = _VALIGN.get(str(p.get("draw:textarea-vertical-align", "top")), MSO_ANCHOR.TOP)
        self.fill_text_frame(tf, item.paragraphs)
        rot = as_float(p.get("librevenge:rotate"), 0.0)
        if rot:
            shp.rotation = rot

    def fill_text_frame(self, tf, paragraphs: list[Paragraph], blank_pt: float | None = None) -> None:
        """Write paragraphs into a text frame.

        `blank_pt` caps the font size given to paragraphs that arrived with no span at all
        (libmspub drops the paragraph mark's style inside table cells); see estimate_blank_size.
        """
        if not paragraphs:
            return
        # Empty paragraphs carry no style in the dump; borrow the nearest text's font size
        # so blank spacer lines keep the height they had in Publisher.
        ref: list[dict[str, Any] | None] = [None] * len(paragraphs)
        last = None
        for i in range(len(paragraphs) - 1, -1, -1):
            if paragraphs[i].runs:
                last = paragraphs[i].runs[0].props
            ref[i] = last
        for i, para in enumerate(paragraphs):
            if ref[i] is None:
                ref[i] = last
            if para.runs:
                last = para.runs[-1].props
        first = True
        for para, rp in zip(paragraphs, ref):
            pp = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            self.fill_paragraph(pp, para, rp, blank_pt)

    def fill_paragraph(self, pp, para: Paragraph, ref_props: dict[str, Any] | None = None,
                       blank_pt: float | None = None) -> None:
        pr = para.props
        pp.alignment = _ALIGN.get(str(pr.get("fo:text-align", "left")).lower(), PP_ALIGN.LEFT)
        pPr = pp._p.get_or_add_pPr()
        # libmspub occasionally emits absurd margins (thousands of inches); ignore those.
        ml = length(pr.get("fo:margin-left"))
        if ml is not None and abs(ml) < self.page_w:
            pPr.set("marL", str(emu(max(ml, 0))))
        mr = length(pr.get("fo:margin-right"))
        if mr is not None and 0 <= mr < self.page_w:
            pPr.set("marR", str(emu(mr)))
        ti = length(pr.get("fo:text-indent"))
        if ti is not None and abs(ti) < self.page_w:
            pPr.set("indent", str(emu(ti)))
            if ti < 0 and "marL" not in pPr.attrib:
                # hanging indent needs a left margin at least as big
                pPr.set("marL", str(emu(-ti)))
        mt, mb = length(pr.get("fo:margin-top")), length(pr.get("fo:margin-bottom"))
        if mt is not None and 0 <= mt < 2:
            pp.space_before = Pt(mt * 72)
        if mb is not None and 0 <= mb < 2:
            pp.space_after = Pt(mb * 72)
        lh = pr.get("fo:line-height")
        if lh is not None:
            s = str(lh).strip()
            if s.endswith("%"):
                pp.line_spacing = as_float(s[:-1], 100) / 100
            else:
                lv = length(s)
                if lv and lv > 0:
                    pp.line_spacing = Pt(lv * 72)
        if para.bullet:
            for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
                for el in pPr.findall(qn(tag)):
                    pPr.remove(el)
            pPr.append(pPr.makeelement(qn("a:buChar"), {"char": para.bullet}))

        for run in para.runs:
            pieces = run.text.split("\n")
            for i, piece in enumerate(pieces):
                if i > 0:
                    pp.add_line_break()
                if piece == "":
                    continue
                r = pp.add_run()
                r.text = piece
                self.style_run(r, run.props)
        # Keep the paragraph's size even if it has no text (blank lines matter for layout).
        if not pp.runs:
            if para.runs:
                # A blank paragraph that kept its own span (text boxes): the size is exact.
                self.style_end_para(pp, para.runs[0].props)
            elif ref_props or blank_pt:
                # No span at all (table cells): borrow the neighbour's style, capped by the estimate.
                size_pt = length((ref_props or {}).get("fo:font-size"))
                size_pt = size_pt * 72 if size_pt else None
                if blank_pt is not None:
                    size_pt = min(size_pt, blank_pt) if size_pt else blank_pt
                self.style_end_para(pp, ref_props or {}, size_pt)

    def style_end_para(self, pp, props: dict[str, Any], size_pt: float | None = None) -> None:
        end = pp._p.find(qn("a:endParaRPr"))
        if end is None:
            end = pp._p.makeelement(qn("a:endParaRPr"), {})
            pp._p.append(end)
        if size_pt is None:
            fs = length(props.get("fo:font-size"))
            size_pt = fs * 72 if fs else None
        if size_pt:
            end.set("sz", str(int(round(max(size_pt, 1) * 100))))
        name = props.get("style:font-name")
        if name and end.find(qn("a:latin")) is None:
            end.append(end.makeelement(qn("a:latin"), {"typeface": self.font_map.get(name, name)}))

    def style_run(self, r, props: dict[str, Any]) -> None:
        f = r.font
        name = props.get("style:font-name")
        if name:
            f.name = self.font_map.get(name, name)
        fs = length(props.get("fo:font-size"))
        if fs:
            f.size = Pt(max(fs * 72, 1))
        if str(props.get("fo:font-weight", "")).lower() in ("bold", "700", "800", "900"):
            f.bold = True
        if str(props.get("fo:font-style", "")).lower() in ("italic", "oblique"):
            f.italic = True
        ul = str(props.get("style:text-underline-type", props.get("style:text-underline-style", "none"))).lower()
        if ul not in ("", "none"):
            f.underline = True
        color = rgb(props.get("fo:color"))
        if color:
            f.color.rgb = color
        rPr = r._r.get_or_add_rPr()
        if str(props.get("fo:font-variant", "")).lower() == "small-caps":
            rPr.set("cap", "small")
        if str(props.get("fo:text-transform", "")).lower() == "uppercase":
            rPr.set("cap", "all")
        lt = str(props.get("style:text-line-through-type", props.get("style:text-line-through-style", "none"))).lower()
        if lt not in ("", "none"):
            rPr.set("strike", "sngStrike")
        pos = str(props.get("style:text-position", "")).strip().lower()
        if pos:
            first = pos.split()[0]
            if first.startswith("super"):
                rPr.set("baseline", "30000")
            elif first.startswith("sub"):
                rPr.set("baseline", "-25000")
            elif first.endswith("%"):
                v = as_float(first[:-1], 0)
                if v:
                    rPr.set("baseline", str(int(v * 1000)))
        lang = props.get("fo:language")
        country = props.get("fo:country")
        if lang:
            rPr.set("lang", f"{lang}-{country}" if country else str(lang))

    # ---- tables ----------------------------------------------------------- #
    LINE_FACTOR = 1.2  # typical line height as a multiple of the font size

    @staticmethod
    def _para_size_pt(para: Paragraph, fallback: float) -> float:
        sizes = [length(r.props.get("fo:font-size")) for r in para.runs]
        sizes = [x * 72 for x in sizes if x]
        return max(sizes) if sizes else fallback

    def _para_height_pt(self, para: Paragraph, avail_width_pt: float, fallback_pt: float) -> float:
        """Rough height of a paragraph with text: lines * line height + spacing."""
        size = self._para_size_pt(para, fallback_pt)
        text = "".join(r.text for r in para.runs)
        lines = 0
        for chunk in text.split("\n"):
            width = len(chunk) * size * 0.5  # average glyph ~0.5em
            lines += max(1, math.ceil(width / max(avail_width_pt, size)))
        factor = self.LINE_FACTOR
        lh = str(para.props.get("fo:line-height", "")).strip()
        if lh.endswith("%"):
            factor *= as_float(lh[:-1], 100) / 100
        elif lh:
            lv = length(lh)
            if lv:
                factor = lv * 72 / size
        h = lines * size * factor
        for key in ("fo:margin-top", "fo:margin-bottom"):
            v = length(para.props.get(key))
            if v and 0 <= v < 2:
                h += v * 72
        return h

    def _cell_geometry(self, cd: Cell, col_w: list[float], n_cols: int):
        """(available text width in pt, vertical padding in inches) for a cell."""
        cs = int(as_float(cd.props.get("table:number-columns-spanned"), 1))
        width_in = sum(col_w[cd.col:min(cd.col + cs, n_cols)])
        pad_v = length(cd.props.get("fo:padding-top"), 0.04) + length(cd.props.get("fo:padding-bottom"), 0.04)
        avail_w = (width_in - length(cd.props.get("fo:padding-left"), 0.04)
                   - length(cd.props.get("fo:padding-right"), 0.04)) * 72
        return max(avail_w, 1.0), pad_v

    @staticmethod
    def _table_min_size(item: Table, default: float = 10.0) -> float:
        sizes = [length(r.props.get("fo:font-size")) for row in item.rows for c in row.cells
                 for p in c.paragraphs for r in p.runs]
        sizes = [x * 72 for x in sizes if x]
        return min(sizes) if sizes else default

    def estimate_blank_size(self, item: Table, col_w: list[float], row_h: list[float], n_cols: int,
                            table_h: float) -> float | None:
        """
        Estimate the font size of blank spacer lines in a table's cells.

        libmspub drops the paragraph mark's style for blank paragraphs inside table cells,
        so their size is unknown. Publisher keeps a row at its stored height unless the
        content overflows, in which case the table's total height grows past the row sum.
        So when the table did not grow, every cell's content fits its row and
        (row height - visible text height) / number of blank lines bounds the blank size.
        The tightest bound across the table is the best available estimate.
        """
        if self.blank_size is not None:
            return self.blank_size
        if table_h - sum(row_h) > 0.02:
            return None  # some rows grew to fit text; the stored heights say nothing
        best: float | None = None
        for r_idx, row in enumerate(item.rows):
            if r_idx >= len(row_h):
                continue
            for cd in row.cells:
                blanks = [p for p in cd.paragraphs if not p.runs]
                texted = [p for p in cd.paragraphs if p.runs]
                if not blanks or not texted:
                    continue
                rs = int(as_float(cd.props.get("table:number-rows-spanned"), 1))
                height_in = sum(row_h[r_idx:min(r_idx + rs, len(row_h))])
                avail_w, pad_v = self._cell_geometry(cd, col_w, n_cols)
                fallback = self._para_size_pt(texted[0], 12.0)
                used = sum(self._para_height_pt(p, avail_w, fallback) for p in texted)
                remaining = (height_in - pad_v) * 72 - used
                if remaining <= 0:
                    continue  # the visible text alone overflows; no information here
                bound = max(4.0, min(remaining / (len(blanks) * self.LINE_FACTOR), fallback))
                best = bound if best is None else min(best, bound)
        if best is not None:
            best = round(best * 2) / 2
            vlog(f"table at ({item.props.get('svg:x')}, {item.props.get('svg:y')}): blank lines estimated at {best}pt")
        return best

    def grow_rows(self, item: Table, col_w: list[float], row_h: list[float], n_cols: int,
                  table_h: float, blank_pt: float) -> list[float]:
        """
        Publisher stores nominal row heights; rows that grew to fit their text only show up
        as the table being taller than the row sum. Hand that extra height to the rows whose
        content overflows, in proportion to how much they overflow.
        """
        extra = table_h - sum(row_h)
        if extra <= 0.02 or not row_h:
            return row_h
        need = [0.0] * len(row_h)
        for r_idx, row in enumerate(item.rows[:len(row_h)]):
            content = 0.0
            for cd in row.cells:
                if int(as_float(cd.props.get("table:number-rows-spanned"), 1)) > 1:
                    continue
                avail_w, pad_v = self._cell_geometry(cd, col_w, n_cols)
                fallback = self._para_size_pt(next((p for p in cd.paragraphs if p.runs), Paragraph({})), blank_pt)
                hpt = 0.0
                for para in cd.paragraphs:
                    if para.runs:
                        hpt += self._para_height_pt(para, avail_w, fallback)
                    else:
                        hpt += blank_pt * self.LINE_FACTOR
                content = max(content, pad_v + hpt / 72)
            need[r_idx] = max(0.0, content - row_h[r_idx])
        total = sum(need)
        if total > 0:
            grown = [rh + extra * n / total for rh, n in zip(row_h, need)]
        else:
            grown = [rh + extra / len(row_h) for rh in row_h]
        vlog(f"table at ({item.props.get('svg:x')}, {item.props.get('svg:y')}): "
             f"{extra:.2f}in of growth given to rows {[i for i, n in enumerate(need) if n > 0] or 'all'}")
        return grown

    def table_layout(self, item: Table):
        """
        Resolve a table's geometry: (x, y, w, h, col_w, row_h, n_cols, blank_pt, table_min),
        all lengths in inches, or None for an empty table. Shared by the PPTX writer and
        the thumbnail painter so both agree on row growth and blank-line sizes.
        """
        p = item.props
        cols_def = p.get("librevenge:table-columns") or []
        n_rows = len(item.rows)
        n_cols = len(cols_def)
        for row in item.rows:
            for c in row.cells:
                span = int(as_float(c.props.get("table:number-columns-spanned"), 1))
                n_cols = max(n_cols, c.col + span)
        if n_rows == 0 or n_cols == 0:
            return None
        x, y = length(p.get("svg:x"), 0.0), length(p.get("svg:y"), 0.0)
        col_w = [length(c.get("style:column-width"), 0.0) or 0.0 for c in cols_def]
        col_w += [0.0] * (n_cols - len(col_w))
        w = length(p.get("svg:width")) or sum(col_w) or 1.0
        if sum(col_w) <= 0:
            col_w = [w / n_cols] * n_cols
        row_h = [length(r.props.get("librevenge:row-height"), 0.0) or 0.0 for r in item.rows]
        h = length(p.get("svg:height")) or sum(row_h) or 1.0
        if sum(row_h) <= 0:
            row_h = [h / n_rows] * n_rows
        table_min = self._table_min_size(item)
        blank_pt = self.estimate_blank_size(item, col_w, row_h, n_cols, h)
        row_h = self.grow_rows(item, col_w, row_h, n_cols, h, blank_pt or table_min)
        return x, y, w, h, col_w, row_h, n_cols, blank_pt, table_min

    def render_table(self, slide, item: Table) -> None:
        layout = self.table_layout(item)
        if layout is None:
            return
        x, y, w, h, col_w, row_h, n_cols, blank_pt, table_min = layout
        n_rows = len(item.rows)
        gf = slide.shapes.add_table(n_rows, n_cols, emu(x), emu(y), emu(w), emu(h))
        tbl = gf.table
        tbl.first_row = False
        tbl.horz_banding = False
        style_id = tbl._tbl.tblPr.find(qn("a:tableStyleId"))
        if style_id is None:
            style_id = tbl._tbl.tblPr.makeelement(qn("a:tableStyleId"), {})
            tbl._tbl.tblPr.append(style_id)
        style_id.text = NO_STYLE_NO_GRID
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = emu(max(cw, 0.01))
        for i, rh in enumerate(row_h):
            tbl.rows[i].height = emu(max(rh, 0.01))

        # Defaults: transparent cells, no borders, small margins.
        for r in range(n_rows):
            for c in range(n_cols):
                cell = tbl.cell(r, c)
                cell.fill.background()
                cell.margin_left = cell.margin_right = emu(0.04)
                cell.margin_top = cell.margin_bottom = emu(0.02)
                self._cell_borders(cell, {})

        merges: list[tuple[int, int, int, int]] = []
        for row in item.rows:
            for cd in row.cells:
                if cd.row >= n_rows or cd.col >= n_cols:
                    continue
                cell = tbl.cell(cd.row, cd.col)
                cp = cd.props
                bg = rgb(cp.get("fo:background-color"))
                if bg:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = bg
                for side in ("left", "right", "top", "bottom"):
                    v = length(cp.get(f"fo:padding-{side}", cp.get("fo:padding")))
                    if v is not None:
                        setattr(cell, f"margin_{side}", emu(v))
                va = str(cp.get("style:vertical-align", cp.get("draw:textarea-vertical-align", "top"))).lower()
                cell.vertical_anchor = _VALIGN.get(va, MSO_ANCHOR.TOP)
                self._cell_borders(cell, cp)
                if any(p.runs for p in cd.paragraphs):
                    self.fill_text_frame(cell.text_frame, cd.paragraphs, blank_pt)
                else:
                    # Nothing but blank lines: keep them small so they never push the row taller.
                    self.fill_text_frame(cell.text_frame, cd.paragraphs, min(table_min, blank_pt or table_min))
                cs = int(as_float(cp.get("table:number-columns-spanned"), 1))
                rs = int(as_float(cp.get("table:number-rows-spanned"), 1))
                if cs > 1 or rs > 1:
                    merges.append((cd.row, cd.col, min(cd.row + rs - 1, n_rows - 1), min(cd.col + cs - 1, n_cols - 1)))
        for r0, c0, r1, c1 in merges:
            try:
                tbl.cell(r0, c0).merge(tbl.cell(r1, c1))
            except Exception as e:
                vlog(f"merge failed ({r0},{c0})-({r1},{c1}): {e}")

    def _cell_borders(self, cell, cp: dict[str, Any]) -> None:
        tcPr = cell._tc.get_or_add_tcPr()
        for side, tag in (("left", "a:lnL"), ("right", "a:lnR"), ("top", "a:lnT"), ("bottom", "a:lnB")):
            for el in tcPr.findall(qn(tag)):
                tcPr.remove(el)
        # Border elements must precede the fill element inside tcPr.
        fill_el = None
        for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill"):
            fill_el = tcPr.find(qn(tag))
            if fill_el is not None:
                break
        for side, tag in (("left", "a:lnL"), ("right", "a:lnR"), ("top", "a:lnT"), ("bottom", "a:lnB")):
            spec = str(cp.get(f"fo:border-{side}", cp.get("fo:border", "none"))).strip()
            ln = tcPr.makeelement(qn(tag), {})
            parts = spec.split()
            color = next((rgb(t) for t in parts if t.startswith("#")), None)
            width = next((length(t) for t in parts if _LEN_RE.match(t)), None)
            if spec.lower() == "none" or not spec or "none" in parts or not color:
                ln.set("w", "0")
                ln.append(ln.makeelement(qn("a:noFill"), {}))
            else:
                ln.set("w", str(emu(width or 0.01)))
                sf = ln.makeelement(qn("a:solidFill"), {})
                sf.append(sf.makeelement(qn("a:srgbClr"), {"val": str(color)}))
                ln.append(sf)
            if fill_el is not None:
                fill_el.addprevious(ln)
            else:
                tcPr.append(ln)


# --------------------------------------------------------------------------- #
# Thumbnail painter (Pillow): draws a page from the document model
# --------------------------------------------------------------------------- #

def _rgb_tuple(v: Any, default=(0, 0, 0)) -> tuple[int, int, int]:
    c = rgb(v)
    return (c[0], c[1], c[2]) if c else default


_FONT_DIRS = [
    "/usr/share/fonts", "/usr/local/share/fonts", "~/.local/share/fonts", "~/.fonts",
    "/Library/Fonts", "/System/Library/Fonts", "~/Library/Fonts",
    os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts"),
] + ([os.path.join(os.environ["LOCALAPPDATA"], "Microsoft", "Windows", "Fonts")]
     if os.environ.get("LOCALAPPDATA") else [])
_FALLBACK_FAMILIES = ["DejaVuSans", "LiberationSans", "NotoSans", "Arial", "FreeSans", "Helvetica"]
_font_index: dict[str, str] | None = None


def _system_fonts() -> dict[str, str]:
    """Lower-cased file stem -> path for every TTF/OTF under the usual font directories."""
    global _font_index
    if _font_index is None:
        _font_index = {}
        for d in _FONT_DIRS:
            root = Path(os.path.expanduser(d))
            if not root.is_dir():
                continue
            for f in root.rglob("*"):
                if f.suffix.lower() in (".ttf", ".otf") and f.is_file():
                    _font_index.setdefault(f.stem.lower(), str(f))
    return _font_index


def _find_font_file(family: str, bold: bool, italic: bool) -> str | None:
    """Locate a font file by name without fontconfig: the family first, then a common sans."""
    index = _system_fonts()
    style_tags = {
        (False, False): ["-regular", "regular", "", "-book"],
        (True, False): ["-bold", "bold", "bd", "-semibold", "b"],
        (False, True): ["-italic", "italic", "-oblique", "i", "-regularitalic"],
        (True, True): ["-bolditalic", "bolditalic", "-boldoblique", "bi", "z"],
    }
    families = [family.replace(" ", ""), family] if family else []
    families += _FALLBACK_FAMILIES
    for fam in families:
        base = fam.lower()
        for tag in style_tags[(bold, italic)] + style_tags[(False, False)]:
            for stem in (base + tag, base + " " + tag.strip("-")):
                path = index.get(stem.strip())
                if path:
                    return path
    return None


def _fc_match(family: str, bold: bool, italic: bool) -> str | None:
    """A font file for the family/style: fontconfig when available, else a directory search."""
    fc = shutil.which("fc-match")
    if not fc:
        return _find_font_file(family, bold, italic)
    pattern = f"{family}:weight={'bold' if bold else 'regular'}:slant={'italic' if italic else 'roman'}"
    try:
        out = subprocess.run([fc, "-f", "%{file}", pattern], capture_output=True, timeout=10)
        path = out.stdout.decode(errors="replace").strip()
        if path and os.path.exists(path):
            return path
    except Exception:
        pass
    return _find_font_file(family, bold, italic)


class ThumbnailPainter:
    """
    Paints one page with Pillow using the same document model the PPTX is built from.
    It is a preview, not a renderer: per-paragraph styling, simple word wrapping,
    curves flattened, gradients reduced to a colour. Good enough for a 256px thumbnail
    and it needs nothing beyond Pillow (fontconfig's fc-match is used when present).
    """

    def __init__(self, renderer: "Renderer", page: Page, max_px: int = 512):
        from PIL import Image, ImageDraw
        self.r = renderer
        self.page = page
        pw = length(page.props.get("svg:width"), renderer.page_w) or renderer.page_w
        ph = length(page.props.get("svg:height"), renderer.page_h) or renderer.page_h
        self.scale = max_px / max(pw, ph)  # pixels per inch
        self.im = Image.new("RGB", (max(1, int(pw * self.scale)), max(1, int(ph * self.scale))), "white")
        self.draw = ImageDraw.Draw(self.im, "RGBA")
        self._fonts: dict[tuple, Any] = {}
        self._font_files: dict[tuple, str | None] = {}

    # ---- helpers ---------------------------------------------------------- #
    def px(self, inches: float) -> int:
        return int(round(inches * self.scale))

    def font(self, family: str | None, size_pt: float, bold: bool, italic: bool):
        from PIL import ImageFont
        size_px = max(1, int(round(size_pt / 72 * self.scale)))
        key = (family, bold, italic, size_px)
        if key in self._fonts:
            return self._fonts[key]
        fkey = (family, bold, italic)
        if fkey not in self._font_files:
            self._font_files[fkey] = _fc_match(self.r.font_map.get(family, family) if family else "sans-serif", bold, italic)
        path = self._font_files[fkey]
        f = None
        if path:
            try:
                f = ImageFont.truetype(path, size_px)
            except Exception:
                f = None
        if f is None:
            f = ImageFont.load_default(size=size_px)
        self._fonts[key] = f
        return f

    def image(self, style: dict[str, Any]):
        from PIL import Image
        data = self.r.image_bytes(style.get("draw:fill-image", ""), style.get("librevenge:mime-type", ""))
        if not data:
            return None
        try:
            im = Image.open(io.BytesIO(data))
            im.load()
            return im.convert("RGBA")
        except Exception:
            return None

    def fill_color(self, style: dict[str, Any]):
        """RGBA fill for a shape style, or None."""
        fill = style.get("draw:fill", "none")
        if fill == "solid":
            a = percent(style.get("draw:opacity"), 1.0)
            return _rgb_tuple(style.get("draw:fill-color")) + (int(255 * min(1.0, max(0.0, a or 1.0))),)
        if fill == "gradient":
            stops = style.get("svg:linearGradient") or style.get("svg:radialGradient") or []
            colors = [rgb(st.get("svg:stop-color")) for st in stops]
            colors = [c for c in colors if c] or [rgb(style.get("draw:start-color")), rgb(style.get("draw:end-color"))]
            colors = [c for c in colors if c]
            if colors:
                a, b = colors[0], colors[-1]
                return ((a[0] + b[0]) // 2, (a[1] + b[1]) // 2, (a[2] + b[2]) // 2, 255)
        return None

    def stroke(self, style: dict[str, Any]):
        """(RGBA colour, width px) for a stroke, or None."""
        if style.get("draw:stroke", "none") == "none":
            return None
        w = length(style.get("svg:stroke-width"), 0.01) or 0.01
        return _rgb_tuple(style.get("svg:stroke-color")) + (255,), max(1, self.px(w))

    def paste_rotated(self, layer, cx_px: float, cy_px: float, angle_deg: float) -> None:
        """Paste an RGBA layer centred at (cx, cy), rotated clockwise by angle_deg."""
        if abs(angle_deg) > 1e-3:
            layer = layer.rotate(-angle_deg, expand=True, resample=3)  # PIL rotates counter-clockwise
        self.im.paste(layer, (int(round(cx_px - layer.width / 2)), int(round(cy_px - layer.height / 2))), layer)

    def picture_layer(self, style: dict[str, Any], w_px: int, h_px: int):
        """The bitmap fill of a style as a w x h RGBA layer (stretched or tiled), or None."""
        from PIL import Image
        src = self.image(style)
        if src is None or w_px < 1 or h_px < 1:
            return None
        if style.get("style:repeat", "repeat") in ("stretch", "no-repeat"):
            return src.resize((w_px, h_px), resample=3)
        tw = max(1, int(round(src.width * self.scale / 96)))  # tiles at 96 dpi native size
        th = max(1, int(round(src.height * self.scale / 96)))
        if (w_px // tw) * (h_px // th) > 4000:  # absurdly small tiles: just stretch
            return src.resize((w_px, h_px), resample=3)
        tile = src.resize((tw, th), resample=3)
        layer = Image.new("RGBA", (w_px, h_px), (0, 0, 0, 0))
        for yy in range(0, h_px, th):
            for xx in range(0, w_px, tw):
                layer.paste(tile, (xx, yy))
        return layer

    # ---- page ------------------------------------------------------------- #
    def paint(self):
        for item in self.page.items:
            try:
                if isinstance(item, Shape):
                    self.paint_shape(item)
                elif isinstance(item, TextBox):
                    self.paint_textbox(item)
                elif isinstance(item, Table):
                    self.paint_table(item)
            except Exception as e:
                vlog(f"thumbnail: skipped {type(item).__name__}: {e}")
        return self.im

    # ---- shapes ----------------------------------------------------------- #
    def paint_shape(self, item: Shape) -> None:
        st, p = item.style, item.props
        rot = as_float(p.get("librevenge:rotate", st.get("librevenge:rotate")), 0.0)
        match item.kind:
            case "graphicobject":
                from PIL import Image
                data = self.r.image_bytes(p.get("office:binary-data", ""), p.get("librevenge:mime-type", ""))
                if not data:
                    return
                im = Image.open(io.BytesIO(data)).convert("RGBA")
                x, y = length(p.get("svg:x"), 0.0), length(p.get("svg:y"), 0.0)
                w, h = length(p.get("svg:width"), 1.0), length(p.get("svg:height"), 1.0)
                layer = im.resize((max(1, self.px(w)), max(1, self.px(h))), resample=3)
                self.paste_rotated(layer, self.px(x + w / 2), self.px(y + h / 2), rot)
            case "polygon":
                pts = points_of(p.get("svg:points"))
                if rect := as_rectangle(pts):
                    cx, cy, w, h, angle = rect
                    self.paint_rect(st, cx, cy, w, h, angle)
                elif pts:
                    self.paint_polygon(st, pts, closed=True)
            case "rectangle":
                x, y = length(p.get("svg:x"), 0.0), length(p.get("svg:y"), 0.0)
                w, h = length(p.get("svg:width"), 0.0), length(p.get("svg:height"), 0.0)
                self.paint_rect(st, x + w / 2, y + h / 2, w, h, rot)
            case "ellipse":
                cx, cy = length(p.get("svg:cx"), 0.0), length(p.get("svg:cy"), 0.0)
                rx, ry = length(p.get("svg:rx"), 0.0), length(p.get("svg:ry"), 0.0)
                box = [self.px(cx - rx), self.px(cy - ry), self.px(cx + rx), self.px(cy + ry)]
                fill = self.fill_color(st)
                stroke = self.stroke(st)
                self.draw.ellipse(box, fill=fill, outline=stroke[0] if stroke else None,
                                  width=stroke[1] if stroke else 1)
            case "polyline" | "connector":
                pts = points_of(p.get("svg:points"))
                if len(pts) >= 2:
                    self.paint_polygon(st, pts, closed=False)
            case "path":
                cmds = path_commands(p.get("svg:d") or [])
                closed = any(c[0] == "Z" for c in cmds)
                for sub in self.flatten(cmds):
                    if len(sub) >= 2:
                        self.paint_polygon(st, sub, closed=closed)

    @staticmethod
    def flatten(cmds, steps: int = 12) -> list[list[tuple[float, float]]]:
        subs: list[list[tuple[float, float]]] = []
        cur = (0.0, 0.0)
        for c in cmds:
            if c[0] == "M":
                subs.append([(c[1], c[2])])
                cur = (c[1], c[2])
            elif c[0] == "L":
                if not subs:
                    subs.append([cur])
                subs[-1].append((c[1], c[2]))
                cur = (c[1], c[2])
            elif c[0] in ("C", "Q"):
                if not subs:
                    subs.append([cur])
                x0, y0 = cur
                for i in range(1, steps + 1):
                    t = i / steps
                    if c[0] == "C":
                        x1, y1, x2, y2, x3, y3 = c[1:]
                        mt = 1 - t
                        px_ = mt ** 3 * x0 + 3 * mt ** 2 * t * x1 + 3 * mt * t ** 2 * x2 + t ** 3 * x3
                        py_ = mt ** 3 * y0 + 3 * mt ** 2 * t * y1 + 3 * mt * t ** 2 * y2 + t ** 3 * y3
                    else:
                        x1, y1, x2, y2 = c[1:]
                        mt = 1 - t
                        px_ = mt ** 2 * x0 + 2 * mt * t * x1 + t ** 2 * x2
                        py_ = mt ** 2 * y0 + 2 * mt * t * y1 + t ** 2 * y2
                    subs[-1].append((px_, py_))
                cur = (c[-2], c[-1])
            elif c[0] == "Z" and subs:
                cur = subs[-1][0]
        return subs

    def paint_rect(self, st, cx, cy, w, h, angle) -> None:
        from PIL import Image, ImageDraw
        if w <= 0 or h <= 0:
            return
        w_px, h_px = max(1, self.px(w)), max(1, self.px(h))
        layer = None
        if st.get("draw:fill") == "bitmap":
            layer = self.picture_layer(st, w_px, h_px)
        if layer is None:
            fill = self.fill_color(st)
            stroke = self.stroke(st)
            if fill is None and stroke is None:
                return
            layer = Image.new("RGBA", (w_px, h_px), (0, 0, 0, 0))
            d = ImageDraw.Draw(layer, "RGBA")
            d.rectangle([0, 0, w_px - 1, h_px - 1], fill=fill, outline=stroke[0] if stroke else None,
                        width=stroke[1] if stroke else 1)
        else:
            stroke = self.stroke(st)
            if stroke:
                ImageDraw.Draw(layer, "RGBA").rectangle([0, 0, w_px - 1, h_px - 1], outline=stroke[0], width=stroke[1])
        self.paste_rotated(layer, self.px(cx), self.px(cy), angle)

    def paint_polygon(self, st, pts, closed: bool) -> None:
        from PIL import Image
        xy = [(self.px(x), self.px(y)) for x, y in pts]
        stroke = self.stroke(st)
        if closed and len(xy) >= 3:
            if st.get("draw:fill") == "bitmap":
                xs, ys = [q[0] for q in xy], [q[1] for q in xy]
                x0, y0, x1, y1 = min(xs), min(ys), max(xs), max(ys)
                layer = self.picture_layer(st, max(1, x1 - x0), max(1, y1 - y0))
                if layer is not None:
                    mask = Image.new("L", self.im.size, 0)
                    from PIL import ImageDraw
                    ImageDraw.Draw(mask).polygon(xy, fill=255)
                    full = Image.new("RGBA", self.im.size, (0, 0, 0, 0))
                    full.paste(layer, (x0, y0))
                    self.im.paste(full, (0, 0), Image.composite(full.getchannel("A"), mask, mask))
            else:
                fill = self.fill_color(st)
                if fill:
                    self.draw.polygon(xy, fill=fill)
            if stroke:
                self.draw.line(xy + [xy[0]], fill=stroke[0], width=stroke[1])
        elif stroke and len(xy) >= 2:
            self.draw.line(xy, fill=stroke[0], width=stroke[1])

    # ---- text ------------------------------------------------------------- #
    def paint_textbox(self, item: TextBox) -> None:
        p = item.props
        x, y = length(p.get("svg:x"), 0.0), length(p.get("svg:y"), 0.0)
        w, h = max(length(p.get("svg:width"), 1.0), 0.05), max(length(p.get("svg:height"), 0.5), 0.05)
        rot = as_float(p.get("librevenge:rotate"), 0.0)
        if item.style.get("draw:fill", "none") != "none" or item.style.get("draw:stroke", "none") != "none":
            self.paint_rect(item.style, x + w / 2, y + h / 2, w, h, rot)
        pad = (length(p.get("fo:padding-left"), 0.04), length(p.get("fo:padding-top"), 0.04),
               length(p.get("fo:padding-right"), 0.04), length(p.get("fo:padding-bottom"), 0.04))
        valign = str(p.get("draw:textarea-vertical-align", "top"))
        self.paint_text_block(x, y, w, h, item.paragraphs, pad, valign, None, rot)

    def _para_style(self, para: Paragraph, blank_pt: float | None, ref: dict[str, Any] | None):
        """(family, size_pt, bold, italic, colour) for a paragraph, styled by its dominant run."""
        props = ref or {}
        if para.runs:
            props = max(para.runs, key=lambda r: len(r.text)).props or para.runs[0].props
        family = props.get("style:font-name")
        size = length(props.get("fo:font-size"))
        size_pt = size * 72 if size else 12.0
        if not para.runs and blank_pt is not None:
            size_pt = min(size_pt, blank_pt)
        if para.runs:
            sizes = [length(r.props.get("fo:font-size")) for r in para.runs]
            sizes = [v * 72 for v in sizes if v]
            if sizes:
                size_pt = max(sizes)
        bold = str(props.get("fo:font-weight", "")).lower() == "bold"
        italic = str(props.get("fo:font-style", "")).lower() in ("italic", "oblique")
        return family, size_pt, bold, italic, _rgb_tuple(props.get("fo:color"))

    def paint_text_block(self, x, y, w, h, paragraphs: list[Paragraph], pad, valign: str,
                         blank_pt: float | None, rot: float = 0.0) -> None:
        from PIL import Image, ImageDraw
        if not paragraphs:
            return
        w_px, h_px = max(1, self.px(w)), max(1, self.px(h))
        layer = Image.new("RGBA", (w_px, h_px), (0, 0, 0, 0))
        d = ImageDraw.Draw(layer, "RGBA")
        left, top, right, bottom = (self.px(v) for v in pad)
        avail = max(1, w_px - left - right)

        # Blank paragraphs borrow the nearest text's style (same rule as the PPTX writer).
        ref: list[dict[str, Any] | None] = [None] * len(paragraphs)
        last = None
        for i in range(len(paragraphs) - 1, -1, -1):
            if paragraphs[i].runs:
                last = paragraphs[i].runs[0].props
            ref[i] = last
        for i, para in enumerate(paragraphs):
            if ref[i] is None:
                ref[i] = last
            if para.runs:
                last = para.runs[-1].props

        lines: list[tuple[str, Any, tuple, int, str, int]] = []  # text, font, colour, line height, align, indent
        for para, rp in zip(paragraphs, ref):
            family, size_pt, bold, italic, color = self._para_style(para, blank_pt, rp)
            font = self.font(family, size_pt, bold, italic)
            lh = size_pt / 72 * self.scale * 1.2
            lhp = str(para.props.get("fo:line-height", "")).strip()
            if lhp.endswith("%"):
                lh *= as_float(lhp[:-1], 100) / 100
            align = str(para.props.get("fo:text-align", "left")).lower()
            ml = length(para.props.get("fo:margin-left"))
            indent = self.px(ml) if ml is not None and 0 <= ml < self.r.page_w else 0
            text = "".join(r.text for r in para.runs).replace("\t", "    ")
            mt = length(para.props.get("fo:margin-top"))
            if mt and 0 <= mt < 2:
                lines.append(("", font, color, self.px(mt), align, indent))
            for chunk in text.split("\n"):
                words = chunk.split(" ")
                cur = ""
                for word in words:
                    trial = word if not cur else cur + " " + word
                    if font.getlength(trial) <= avail - indent or not cur:
                        cur = trial
                    else:
                        lines.append((cur, font, color, int(lh), align, indent))
                        cur = word
                lines.append((cur, font, color, int(lh), align, indent))
            mb = length(para.props.get("fo:margin-bottom"))
            if mb and 0 <= mb < 2:
                lines.append(("", font, color, self.px(mb), align, indent))

        total = sum(l[3] for l in lines)
        yy = top
        if valign == "middle":
            yy = max(top, (h_px - total) // 2)
        elif valign == "bottom":
            yy = max(top, h_px - bottom - total)
        for text, font, color, lh, align, indent in lines:
            if text:
                tw = font.getlength(text)
                if align == "center":
                    xx = left + indent + (avail - indent - tw) / 2
                elif align in ("right", "end"):
                    xx = w_px - right - tw
                else:
                    xx = left + indent
                d.text((xx, yy), text, font=font, fill=color + (255,))
            yy += lh
        self.paste_rotated(layer, self.px(x + w / 2), self.px(y + h / 2), rot)

    # ---- tables ----------------------------------------------------------- #
    def paint_table(self, item: Table) -> None:
        layout = self.r.table_layout(item)
        if layout is None:
            return
        x, y, w, h, col_w, row_h, n_cols, blank_pt, table_min = layout
        col_x = [x]
        for cw in col_w:
            col_x.append(col_x[-1] + cw)
        row_y = [y]
        for rh in row_h:
            row_y.append(row_y[-1] + rh)
        for r_idx, row in enumerate(item.rows):
            if r_idx >= len(row_h):
                break
            for cd in row.cells:
                if cd.col >= n_cols:
                    continue
                cs = int(as_float(cd.props.get("table:number-columns-spanned"), 1))
                rs = int(as_float(cd.props.get("table:number-rows-spanned"), 1))
                cx0, cy0 = col_x[cd.col], row_y[r_idx]
                cw = col_x[min(cd.col + cs, n_cols)] - cx0
                ch = row_y[min(r_idx + rs, len(row_h))] - cy0
                bg = rgb(cd.props.get("fo:background-color"))
                if bg:
                    self.draw.rectangle([self.px(cx0), self.px(cy0), self.px(cx0 + cw), self.px(cy0 + ch)],
                                        fill=(bg[0], bg[1], bg[2], 255))
                for side, (ax, ay, bx, by) in (("top", (cx0, cy0, cx0 + cw, cy0)), ("bottom", (cx0, cy0 + ch, cx0 + cw, cy0 + ch)),
                                               ("left", (cx0, cy0, cx0, cy0 + ch)), ("right", (cx0 + cw, cy0, cx0 + cw, cy0 + ch))):
                    spec = str(cd.props.get(f"fo:border-{side}", cd.props.get("fo:border", "none"))).split()
                    color = next((rgb(t) for t in spec if t.startswith("#")), None)
                    if color and "none" not in spec:
                        width = next((length(t) for t in spec if _LEN_RE.match(t)), None) or 0.01
                        self.draw.line([(self.px(ax), self.px(ay)), (self.px(bx), self.px(by))],
                                       fill=(color[0], color[1], color[2], 255), width=max(1, self.px(width)))
                pad = (length(cd.props.get("fo:padding-left"), 0.04), length(cd.props.get("fo:padding-top"), 0.02),
                       length(cd.props.get("fo:padding-right"), 0.04), length(cd.props.get("fo:padding-bottom"), 0.02))
                valign = str(cd.props.get("style:vertical-align", "top")).lower()
                has_text = any(p.runs for p in cd.paragraphs)
                bp = blank_pt if has_text else min(table_min, blank_pt or table_min)
                self.paint_text_block(cx0, cy0, cw, ch, cd.paragraphs, pad, valign, bp)


def make_thumbnail_pillow(renderer: "Renderer", doc: Document) -> bytes | None:
    """First-page preview painted with Pillow, encoded as a small JPEG."""
    if not doc.pages:
        return None
    im = ThumbnailPainter(renderer, doc.pages[0]).paint()
    im.thumbnail((THUMBNAIL_PX, THUMBNAIL_PX), resample=3)
    buf = io.BytesIO()
    im.convert("RGB").save(buf, "JPEG", quality=85, optimize=True)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# Drivers
# --------------------------------------------------------------------------- #

_PUB2RAW_HINT = ("install libmspub: Arch `pacman -S libmspub`, Debian/Ubuntu `apt install libmspub-tools`, "
                 "macOS `brew install libmspub`, Windows (MSYS2) `pacman -S mingw-w64-ucrt-x86_64-libmspub`; "
                 "or set PUB2RAW to the executable")


def find_pub2raw() -> str | None:
    """pub2raw from $PUB2RAW, the PATH, or the usual MSYS2 folders on Windows."""
    env = os.environ.get("PUB2RAW")
    if env and os.path.isfile(env):
        return env
    exe = shutil.which("pub2raw")
    if exe:
        return exe
    if sys.platform == "win32":
        for root in (r"C:\msys64", r"C:\msys32", os.path.expanduser("~\\scoop\\apps\\msys2\\current")):
            for sub in ("ucrt64", "mingw64", "clang64", "mingw32"):
                cand = os.path.join(root, sub, "bin", "pub2raw.exe")
                if os.path.isfile(cand):
                    return cand
    return None


def run_pub2raw(pub: Path) -> str:
    exe = find_pub2raw()
    if not exe:
        raise RuntimeError(f"pub2raw not found ({_PUB2RAW_HINT})")
    src = pub.resolve()
    tmp: tempfile.TemporaryDirectory | None = None
    if sys.platform == "win32" and not str(src).isascii():
        # The MinGW build takes its command line in the ANSI code page; sidestep that
        # for names it cannot represent by working on an ASCII-named copy.
        tmp = tempfile.TemporaryDirectory()
        src = Path(tmp.name) / "input.pub"
        shutil.copyfile(pub, src)
    try:
        res = subprocess.run([exe, str(src)], capture_output=True, timeout=600)
    finally:
        if tmp is not None:
            tmp.cleanup()
    if res.returncode != 0 and not res.stdout:
        raise RuntimeError(f"pub2raw failed: {res.stderr.decode(errors='replace').strip()}")
    raw = res.stdout.decode("utf-8", errors="replace")
    # A Windows build writes stdout in text mode, turning every "\n" into "\r\n"; undo that
    # globally so a "\r" left in the stream is a real Publisher line break, as on other platforms.
    if raw.startswith("startDocument()\r\n"):
        raw = raw.replace("\r\n", "\n")
    return raw


THUMBNAIL_PX = 256


def _thumbnail_rel(package):
    """The package's thumbnail relationship (rId, rel), or (None, None)."""
    for rId, rel in package._rels.items():
        if rel.reltype == RT.THUMBNAIL:
            return rId, rel
    return None, None


def embed_thumbnail(prs: Presentation, jpeg: bytes | None) -> None:
    """
    Store the JPEG as docProps/thumbnail.jpeg, linked from the package with the
    standard metadata/thumbnail relationship. PowerPoint, GNOME (via libgsf) and
    KDE (kio-extras office thumbnailer) all read this for file-manager previews.

    python-pptx's default template already carries a blank placeholder thumbnail:
    it is replaced, or removed when `jpeg` is None so no misleading preview remains.
    """
    package = prs.part.package
    rId, rel = _thumbnail_rel(package)
    if jpeg is None:
        if rId is not None:
            package._rels.pop(rId)
        return
    if rel is not None:
        rel.target_part._blob = jpeg
    else:
        part = Part(PackURI("/docProps/thumbnail.jpeg"), CT.JPEG, package, jpeg)
        package.relate_to(part, RT.THUMBNAIL)


def load_document(pub: Path, out: Path, dump_raw: bool) -> Document:
    raw = run_pub2raw(pub)
    if dump_raw:
        raw_path = out.with_suffix(".raw.txt")
        raw_path.write_text(raw, encoding="utf-8")
        vlog(f"wrote {raw_path}")
    doc = build_document(parse_raw(raw))
    if not doc.pages:
        raise RuntimeError("no pages found in document")
    return doc


def convert_editable(pub: Path, out: Path, font_map: dict[str, str], dump_raw: bool,
                     blank_size: float | None = None, thumbnail: bool = True) -> list[str]:
    doc = load_document(pub, out, dump_raw)
    renderer = Renderer(doc, font_map, blank_size)
    prs = renderer.render()
    set_core_props(prs, pub, doc.meta)
    jpeg = None
    if thumbnail:
        try:
            jpeg = make_thumbnail_pillow(renderer, doc)
        except Exception as e:
            log(f"warning: could not paint a thumbnail: {e}")
            if VERBOSE:
                traceback.print_exc()
    embed_thumbnail(prs, jpeg)
    prs.save(str(out))
    return renderer.warnings


def convert_snapshot(pub: Path, out: Path, dpi: int, font_map: dict[str, str], dump_raw: bool,
                     blank_size: float | None = None, thumbnail: bool = True) -> list[str]:
    """Every page painted with Pillow and placed as one full-slide picture."""
    doc = load_document(pub, out, dump_raw)
    renderer = Renderer(doc, font_map, blank_size)
    prs = renderer.prs
    blank = renderer._blank
    for i, page in enumerate(doc.pages):
        pw = length(page.props.get("svg:width"), renderer.page_w) or renderer.page_w
        ph = length(page.props.get("svg:height"), renderer.page_h) or renderer.page_h
        im = ThumbnailPainter(renderer, page, max_px=int(dpi * max(pw, ph))).paint()
        buf = io.BytesIO()
        im.save(buf, "PNG", optimize=True)
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(io.BytesIO(buf.getvalue()), 0, 0, prs.slide_width, prs.slide_height)
        if i == 0:
            jpeg = None
            if thumbnail:
                small = im.copy()
                small.thumbnail((THUMBNAIL_PX, THUMBNAIL_PX), resample=3)
                tb = io.BytesIO()
                small.convert("RGB").save(tb, "JPEG", quality=85, optimize=True)
                jpeg = tb.getvalue()
            embed_thumbnail(prs, jpeg)
    set_core_props(prs, pub, doc.meta)
    prs.save(str(out))
    return renderer.warnings


def set_core_props(prs: Presentation, pub: Path, meta: dict[str, Any] | None = None) -> None:
    from datetime import UTC, datetime
    cp = prs.core_properties
    meta = meta or {}
    cp.title = pub.stem
    cp.author = str(meta.get("dc:creator") or meta.get("meta:initial-creator") or "")
    cp.last_modified_by = "pub2pptx"
    cp.comments = f"Converted from {pub.name} by pub2pptx"
    cp.revision = 1
    now = datetime.now(UTC).replace(microsecond=0, tzinfo=None)
    cp.created = now
    cp.modified = now


def convert_one(pub: Path, out: Path, mode: str, dpi: int, font_map: dict[str, str], dump_raw: bool,
                blank_size: float | None = None, thumbnail: bool = True) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    if mode == "snapshot":
        convert_snapshot(pub, out, dpi, font_map, dump_raw, blank_size, thumbnail)
    else:
        convert_editable(pub, out, font_map, dump_raw, blank_size, thumbnail)


def parse_font_map(spec: str | None) -> dict[str, str]:
    m: dict[str, str] = {}
    if spec:
        for pair in spec.split(","):
            if "=" in pair:
                a, b = pair.split("=", 1)
                m[a.strip()] = b.strip()
    return m


def main(argv: list[str] | None = None) -> int:
    global VERBOSE
    for stream in (sys.stdout, sys.stderr):
        # Never fail on a file name the console encoding can't show (legacy Windows code pages).
        try:
            stream.reconfigure(errors="replace")
        except Exception:
            pass
    ap = argparse.ArgumentParser(prog="pub2pptx", description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("inputs", nargs="+", type=Path, help="Publisher .pub file(s)")
    ap.add_argument("-o", "--output", type=Path,
                    help="output .pptx for a single input, or output directory for several")
    ap.add_argument("--mode", choices=("editable", "snapshot"), default="editable",
                    help="editable (default): native slides; snapshot: one painted picture per page")
    ap.add_argument("--dpi", type=int, default=150, help="paint resolution for snapshot mode (default 150)")
    ap.add_argument("--font-map", metavar="A=B,C=D",
                    help="substitute fonts in editable mode, e.g. 'Raleway=Calibri,Poppins=Segoe UI'")
    ap.add_argument("--blank-size", type=float, metavar="PT",
                    help="font size in points for blank spacer lines inside table cells "
                         "(Publisher files don't expose it; default: estimated from the row height)")
    ap.add_argument("--no-thumbnail", action="store_true",
                    help="don't embed the first-page preview image that file managers show")
    ap.add_argument("--dump-raw", action="store_true", help="also write the pub2raw dump next to the output (debugging)")
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args(argv)
    VERBOSE = args.verbose

    inputs: list[Path] = []
    for p in args.inputs:
        if not p.exists():
            log(f"error: {p} does not exist")
            return 2
        inputs.append(p)

    out_dir: Path | None = None
    out_file: Path | None = None
    if args.output:
        if len(inputs) > 1 or args.output.is_dir() or args.output.suffix.lower() != ".pptx":
            out_dir = args.output
        else:
            out_file = args.output
    ext = ".pptx"
    font_map = parse_font_map(args.font_map)

    failures = 0
    for pub in inputs:
        if out_file:
            out = out_file
        elif out_dir:
            out = out_dir / (pub.stem + ext)
        else:
            out = pub.with_suffix(ext)
        try:
            convert_one(pub, out, args.mode, args.dpi, font_map, args.dump_raw, args.blank_size,
                        not args.no_thumbnail)
            print(f"{pub} -> {out}")
        except Exception as e:
            failures += 1
            log(f"error: {pub}: {e}")
            if VERBOSE:
                traceback.print_exc()
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
